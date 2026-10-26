import copy
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

GREEN, ORANGE, RED, NAVY = RGBColor(0x3F, 0xBA, 0x46), RGBColor(0xE6, 0x9E, 0x0B), RGBColor(0xE7, 0x41, 0x3E), RGBColor(0x00, 0x5B, 0xAC)


def _rpr_templates(p):
    """문단 안의 굵은 run / 보통 run 서식(rPr)을 하나씩 확보한다."""
    b = n = None
    for r in p.runs:
        rpr = r._r.find(qn('a:rPr'))
        if r.font.bold and b is None:
            b = rpr
        elif not r.font.bold and n is None:
            n = rpr
    return b, n


def _apply_color(rpr_el, color):
    for f in rpr_el.findall(qn('a:solidFill')):
        rpr_el.remove(f)
    sf = rpr_el.makeelement(qn('a:solidFill'), {})
    c = sf.makeelement(qn('a:srgbClr'), {'val': str(color)})
    sf.append(c)
    ln = rpr_el.find(qn('a:ln'))
    if ln is not None:
        ln.addnext(sf)
    else:
        rpr_el.insert(0, sf)


def _add_run(p_el, text, rpr, bold=None, color=None):
    parts = text.split('\n')
    for k, part in enumerate(parts):
        if k > 0:
            br = p_el.makeelement(qn('a:br'), {})
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            _insert_before_end(p_el, br)
        if part == '' and k > 0:
            continue
        r = p_el.makeelement(qn('a:r'), {})
        if rpr is not None:
            nr = copy.deepcopy(rpr)
            if bold is not None:
                nr.set('b', '1' if bold else '0')
            if color is not None:
                _apply_color(nr, color)
            r.append(nr)
        t = r.makeelement(qn('a:t'), {})
        t.text = part
        r.append(t)
        _insert_before_end(p_el, r)


def _insert_before_end(p_el, el):
    end = p_el.find(qn('a:endParaRPr'))
    if end is not None:
        end.addprevious(el)
    else:
        p_el.append(el)


def set_para(p, segs, bold_tpl=None, norm_tpl=None):
    """segs = [(text, bold), ...] — 기존 run을 지우고 굵기별 원래 서식으로 다시 채운다."""
    b, n = _rpr_templates(p)
    b = b if b is not None else bold_tpl
    n = n if n is not None else norm_tpl
    if b is None and n is not None:
        b = n
    if n is None and b is not None:
        n = b
    for r in list(p._p):
        if r.tag in (qn('a:r'), qn('a:br'), qn('a:fld')):
            p._p.remove(r)
    for seg in segs:
        text, bold = seg[0], seg[1]
        color = seg[2] if len(seg) > 2 else None
        tpl = b if bold else n
        _add_run(p._p, text, tpl, bold=bold, color=color)


def set_paras(shape, paras):
    """paras = [ [(text,bold),...], ... ] — 문단 수를 맞추고 각 문단을 채운다."""
    tf = shape.text_frame
    ps = list(tf.paragraphs)
    # 서식 템플릿 (전체에서 확보)
    gb = gn = None
    for p in ps:
        b, n = _rpr_templates(p)
        gb = gb if gb is not None else b
        gn = gn if gn is not None else n
    while len(tf.paragraphs) < len(paras):
        last = tf.paragraphs[-1]._p
        last.addnext(copy.deepcopy(last))
    ps = list(tf.paragraphs)
    for extra in ps[len(paras):]:
        extra._p.getparent().remove(extra._p)
    for p, segs in zip(tf.paragraphs, paras):
        if isinstance(segs, str):
            segs = [(segs, False)]
        set_para(p, segs, gb, gn)


def set_text(shape, text, bold=None):
    """단일 문단 텍스트 교체 (첫 run 서식 유지)."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    rs = p.runs
    if rs:
        bb = rs[0].font.bold if bold is None else bold
        set_para(p, [(text, bool(bb))])
    else:
        p.text = text
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)


def set_cell(cell, text, color=None, bold=None):
    tf = cell.text_frame
    p = tf.paragraphs[0]
    rs = p.runs
    if rs:
        rpr = rs[0]._r.find(qn('a:rPr'))
        bb = rs[0].font.bold if bold is None else bold
        for r in list(p._p):
            if r.tag in (qn('a:r'), qn('a:br')):
                p._p.remove(r)
        _add_run(p._p, str(text), rpr, bold=bb)
    else:
        p.text = str(text)
        if bold is not None:
            p.runs[0].font.bold = bold
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    if color is not None:
        for r in p.runs:
            r.font.color.rgb = color


def fill_table(tbl, rows, start_row=1):
    """rows = [[...], ...] — 필요하면 마지막 행을 복제해 늘리고, 남는 행은 지운다."""
    tr_list = tbl._tbl.tr_lst
    need = start_row + len(rows)
    while len(tbl._tbl.tr_lst) < need:
        src = tbl._tbl.tr_lst[-1]
        src.addnext(copy.deepcopy(src))
    for tr in list(tbl._tbl.tr_lst)[need:]:
        tbl._tbl.remove(tr)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is None:
                continue
            set_cell(tbl.cell(start_row + i, j), v)


def set_row_heights(tbl, h_in):
    for tr in tbl._tbl.tr_lst:
        tr.set('h', str(int(Inches(h_in))))


def set_table_font(tbl, pt):
    for r in tbl.rows:
        for c in r.cells:
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(pt)


def replace_chart(chart, cats, series):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    chart.replace_data(cd)


def move(shape, L=None, T=None, W=None, H=None):
    if L is not None: shape.left = Inches(L)
    if T is not None: shape.top = Inches(T)
    if W is not None: shape.width = Inches(W)
    if H is not None: shape.height = Inches(H)


def delete_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def sanitize_tables(prs):
    """PowerPoint 호환 정리 — 행 끝 extLst를 맨 뒤로, 열/행 ID를 표 안에서 유일하게."""
    A16 = '{http://schemas.microsoft.com/office/drawing/2014/main}'
    n_fix = 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_table:
                continue
            tbl = sh.table._tbl
            for k, gc in enumerate(tbl.tblGrid.gridCol_lst):
                for e in gc.iter(A16 + 'colId'):
                    e.set('val', str(20000 + k))
            for k, tr in enumerate(tbl.tr_lst):
                ext = tr.find(qn('a:extLst'))
                if ext is not None and tr.index(ext) != len(tr) - 1:
                    tr.remove(ext); tr.append(ext); n_fix += 1
                for e in tr.iter(A16 + 'rowId'):
                    e.set('val', str(10000 + k))
    return n_fix


def check_tables(prs):
    """표 구조 점검 — 셀 수 = 열 수, extLst 뒤에 셀 없음, ID 중복 없음."""
    A16 = '{http://schemas.microsoft.com/office/drawing/2014/main}'
    problems = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_table:
                continue
            tbl = sh.table._tbl
            ncol = len(tbl.tblGrid.gridCol_lst)
            cids = [e.get('val') for gc in tbl.tblGrid.gridCol_lst for e in gc.iter(A16 + 'colId')]
            rids = [e.get('val') for tr in tbl.tr_lst for e in tr.iter(A16 + 'rowId')]
            if len(cids) != len(set(cids)) or len(rids) != len(set(rids)):
                problems.append((si, sh.name, 'ID 중복'))
            for tr in tbl.tr_lst:
                kids = list(tr)
                tcs = [c for c in kids if c.tag == qn('a:tc')]
                if len(tcs) != ncol:
                    problems.append((si, sh.name, f'셀 {len(tcs)} ≠ 열 {ncol}'))
                ext = tr.find(qn('a:extLst'))
                if ext is not None and kids.index(ext) != len(kids) - 1:
                    problems.append((si, sh.name, 'extLst 뒤에 셀'))
    return problems


def bump_fonts(prs, table_min=10.0, foot_min=10.0, body_min=11.5, skip_slides=(1,)):
    """가독성 최소 크기 — 표 10pt, 각주(9pt 미만) 10pt, 본문(9~11pt) 11.5pt. 제목·큰 숫자는 그대로."""
    for si, s in enumerate(prs.slides, 1):
        if si in skip_slides:
            continue
        for sh in s.shapes:
            if sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        for pa in c.text_frame.paragraphs:
                            for ru in pa.runs:
                                if ru.font.size and ru.font.size.pt < table_min:
                                    ru.font.size = Pt(table_min)
            elif sh.has_text_frame:
                for pa in sh.text_frame.paragraphs:
                    for ru in pa.runs:
                        if not ru.font.size:
                            continue
                        pt = ru.font.size.pt
                        if pt < 9:
                            ru.font.size = Pt(foot_min)
                        elif pt < 11:
                            ru.font.size = Pt(body_min)
