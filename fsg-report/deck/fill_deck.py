#!/usr/bin/env python3
"""월간보고 덱 자동 채움 — report-data_<기간>.json 의 값으로 전월 확정 덱(베이스)의 숫자·표·차트를 바꾼다.

  python fill_deck.py <베이스.pptx> <report-data_2026-10.json> [-o 출력.pptx]

바꾸는 곳(확정 12장 덱 기준): 2 핵심 요약 카드 · 3 팀 표·KPI 카드·차트 · 4 설치 표 · 5 케이스 분석 · 6 트렌드 비교 · 7 수도권 · 8 병원/반복고장/장애요인 · 10 에러율·운영 건전성.
바꾸지 않는 곳: 9 FCA · 11 만족도 · 12 Next Step, 그리고 판단이 필요한 문장(이슈·대책·분석 코멘트).
→ 판단 문장은 값이 낡았을 수 있어 실행 끝에 「검토 필요」 목록으로 따로 알려 준다. json 에 없는 값(엔진이 아직 안 내는 것)은 건드리지 않고 「건너뜀」 목록에 적는다.
베이스 덱이 이미 같은 달이면(시험용) 값만 새 json 으로 갈아 끼운다.
"""
import argparse, json, os, re, sys
from pptx import Presentation
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ppt_helpers import set_text, set_paras, set_cell, fill_table, replace_chart, GREEN, RED

B, N = True, False
SKIP, REVIEW = [], []
TEAM_KO = {'FS EAST': 'FS East', 'FS WEST': 'FS West', 'FS SOUTH': 'FS South', 'FS CENTRAL': 'FS Central'}
TEAM_SHORT = {'FS EAST': 'East', 'FS WEST': 'West', 'FS SOUTH': 'South', 'FS CENTRAL': 'Central'}


PROVINCE = r'^(서울특별시|경상북도|경상남도|전라남도|전라북도|충청남도|충청북도|강원도|경기도|제주특별자치도|부산광역시|대구광역시|인천광역시|광주광역시|대전광역시|울산광역시)\s*'
def short_acct(a, tight=False):
    a = re.sub(PROVINCE, '', a or '').strip(); a = re.sub(r'^(의료법인|학교법인|재단법인)\s*', '', a)
    return a.replace('대학교', '대') if tight else a
def short_model(m): return re.sub(r'\s*COMPLETE.*$', '', m or '').strip()
def f1(v): return f'{v:.1f}'
def arrow(d): return '▲' if d > 0 else ('▼' if d < 0 else '■')
def sgn(d, nd=1, unit=''): return ('+' if d > 0 else '') + f'{d:.{nd}f}{unit}'
def comma(v): return f'{int(round(v)):,}'
def mnum(label): return int(label.split('-')[1])
def yymm(label): return label[2:]


def run(base, data, out):
    prs = Presentation(base)
    S = [list(s.shapes) for s in prs.slides]
    d = data
    tr = d['kpiTrend']; cur, prv = tr[-1]['kpi'], tr[-2]['kpi']
    label = d['meta']['period']['label']; m = mnum(label); pm_ = mnum(tr[-2]['label'])
    mo, pmo = f'{m}월', f'{pm_}월'
    prev_mo_label = tr[-2]['label']
    ib = d.get('installBlock')

    # ── 2. 핵심 요약
    s = S[1]
    set_text(s[0], f'{mo} 핵심 요약 (Executive Summary)')
    for k in (3, 9, 15): set_text(s[k], f'{pmo} 대비')
    def card(shape, p, c, unit, target, higher_good):
        dlt = c - p; ok = (c >= target) if higher_good else (c <= target)
        set_text(shape, f'{f1(p)} → {f1(c)}{unit} ({arrow(dlt)}{abs(dlt):.1f}{"%p" if unit == "%" else unit} · {"달성" if ok else "미달"})')
    card(s[4], prv['ftfr'], cur['ftfr'], '%', 90, True)
    card(s[10], prv['mttr'], cur['mttr'], '분', 120, False)
    card(s[16], prv['revisitRate'], cur['revisitRate'], '%', 15, False)
    t = s[19].text_frame.text
    set_text(s[19], re.sub(r'예방정비 지수 [\d.]+% → [\d.]+%', f'예방정비 지수 {f1(prv["pmRatio"])}% → {f1(cur["pmRatio"])}%', t))
    REVIEW.append(f'2장 「향후 주요 이슈 & 대책」 — 숫자·원인 문장은 사람이 새로 써야 함(현재 텍스트는 베이스 달 기준): "{s[22].text_frame.text[:60]}…"')

    # ── 3. 서비스 활동
    s = S[2]
    set_text(s[0], f'{mo} 서비스 활동 보고')
    team = d['team']
    units = {}; orgs = {}
    if ib:
        for r in ib['summary']['mat']:
            if r[1] == '소계': units[r[0]] = r[-1]
            elif r[1] in ('신규', '이전설치'): orgs[r[0]] = orgs.get(r[0], 0) + r[2]
    rows3 = []; tot = {'cnt': 0, 'org': 0, 'unit': 0, 'pm': 0, 'as': 0, 'sum': 0}
    for r in team['rows']:
        org = orgs.get(r['team'], r['inst']); un = units.get(r['team'])
        inst_txt = f"{org} / {un}" if un is not None else f"{org}"
        rows3.append([None, None, f"{r['count']}명", inst_txt, comma(r['pm']), comma(r['as_']), comma(r['sum']), f1(r['perHead'])])
        tot['cnt'] += r['count']; tot['org'] += org; tot['unit'] += (un or 0); tot['pm'] += r['pm']; tot['as'] += r['as_']; tot['sum'] += r['sum']
    rows3.append([None, None, f"{tot['cnt']}명", f"{tot['org']} / {tot['unit']}" if ib else f"{tot['org']}", comma(tot['pm']), comma(tot['as']), comma(tot['sum']), f1(tot['sum'] / tot['cnt'])])
    if not ib: SKIP.append('3장 설치 칸 장비 대수 — 설치 정보(캘린더·Serial)가 json 에 없음')
    fill_table(s[1].table, rows3)
    tshort = lambda key: ' · '.join(f"{TEAM_SHORT[t['team']][0]} {f1(t[key])}" for t in cur['teams'])
    dt = cur['ftfr'] - prv['ftfr']
    set_text(s[5], f"{f1(cur['ftfr'])}%  목표 90%↑"); set_text(s[6], tshort('ftfr'))
    set_text(s[7], f"{pmo} {f1(prv['ftfr'])} → {mo} {f1(cur['ftfr'])}  {arrow(dt)}{abs(dt):.1f}%p")
    dm = cur['mttr'] - prv['mttr']
    set_text(s[10], f"{f1(cur['mttr'])}분  목표 120분↓"); set_text(s[11], tshort('mttr'))
    set_text(s[12], f"{pmo} {f1(prv['mttr'])} → {mo} {f1(cur['mttr'])}  {arrow(dm)}{abs(dm):.1f}분")
    cats12 = [yymm(x['label']) for x in tr]
    replace_chart(s[15].chart, cats12, [('FTFR(%)', [x['kpi']['ftfr'] for x in tr]), ('목표 90%', [90.0] * 12)])
    replace_chart(s[16].chart, cats12, [('MTTR(분)', [x['kpi']['mttr'] for x in tr]), ('목표 120분', [120.0] * 12)])
    REVIEW.append(f'3장 하단 요약 3줄(전사 활동량·팀 특이사항·품질) — 문장을 새로 써야 함: "{s[14].text_frame.text[:50]}…"')

    # ── 4. 설치 현황
    s = S[3]
    set_text(s[0], f'{mo} 장비별 설치 현황')
    if ib:
        mat = ib['summary']['mat']; byteam = {}
        for r in mat: byteam.setdefault(r[0], {})[r[1]] = r
        rows4 = []
        for t in ['FS EAST', 'FS WEST', 'FS SOUTH', 'FS CENTRAL']:
            g = byteam[t]; sub = g['소계']
            rows4.append([None, g['신규'][2], g['이전설치'][2]] + [str(x) for x in sub[3:]])
        g = byteam['전체']; tt = g['합계']
        rows4.append([None, g['신규'][2], g['이전설치'][2]] + [str(x) for x in tt[3:]])
        fill_table(s[3].table, rows4)
        n_new, n_mv, n_units = g['신규'][2], g['이전설치'][2], tt[-1]
        t1 = s[1].text_frame.text
        t1 = re.sub(r'\d+월 설치 \d+곳 \(신규 \d+ · 이전설치 \d+\) · 장비 \d+대', f'{m}월 설치 {n_new + n_mv}곳 (신규 {n_new} · 이전설치 {n_mv}) · 장비 {n_units}대', t1)
        t1 = re.sub(r'^\d{4}년 \d+월', f'{label[:4]}년 {m}월', t1)
        set_text(s[1], t1)
    else:
        SKIP.append('4장 설치 표 — json 에 installBlock 없음(캘린더·Serial 필요)')
    REVIEW.append('4장 하단 각주(「모듈·구성 모델 중복 제외(v50) · 캘린더 미등록 N곳…」) — 수동입력 폐지 후 규칙(캘린더 장비 표기·자동 연결)으로 문구를 고칠 것')
    SKIP.append('4장 「주요 기관 설치」·「다음 달 설치 예정」 두 표와 「포인트」 문장 — 엔진이 다음 달 일정을 내지 않아 사람이 확인')

    # ── 5. 케이스 분석
    s = S[4]
    ca = d['caseAnalysis']; mr = ca['monthRows']
    replace_chart(s[2].chart, [yymm(r[0]) for r in mr], [('고장 수리', [r[1] for r in mr]), ('PM', [r[2] for r in mr]), ('Calibration', [r[3] for r in mr])])
    set_text(s[5], f"{f1(cur['pmRatio'])}%  ({mo})")
    pmv, rpv, cal = mr[-1][2], mr[-1][1], mr[-1][3]
    pmp, rpp, calp = mr[-2][2], mr[-2][1], mr[-2][3]
    nxt = f'{m % 12 + 1}월'
    set_paras(s[8], [
        [(f"{mo} 예방정비 지수 {f1(prv['pmRatio'])}% → {f1(cur['pmRatio'])}%", B)],
        [(f"• PM {comma(pmv)}건 ({pmo} 대비 ", N), (f"{arrow(pmv - pmp)}{abs(pmv - pmp)}건", B), (')', N)],
        [(f"• 고장수리 {comma(rpv)}건 ({pmo} 대비 ", N), (f"{arrow(rpv - rpp)}{abs(rpv - rpp)}건", B), (')', N)],
        [(f"• Cal {cal}건({pmo} {calp}건) — {nxt} ", N), ('반락 여부 확인', B)],
    ])
    cr = {r[0]: r for r in ca['corrRows']}
    allr = cr['(전체 평균)']
    def rate(r): return float(r[3]); 
    def prate(r): return None if r[4] == '-' else float(r[4])
    ser = [r for r in ca['corrRows'] if r[0] != '(전체 평균)']
    best = max((r for r in ser if prate(r) is not None), key=lambda r: prate(r) - rate(r), default=None)
    hi = max(ser, key=rate)
    da = rate(allr) - (prate(allr) or 0)
    paras = [[('해석', B)],
             [(f"전체 {f1(prate(allr))}% → {f1(rate(allr))}% (", N), (f"{arrow(da)}{abs(da):.1f}%p", B), (')', N)]]
    if best: paras.append([(f"{best[0]} {f1(prate(best))}% → {f1(rate(best))}% — 개선 폭 최대", N)])
    paras.append([(f"{hi[0]} ", N), (f"{f1(rate(hi))}%", B), (f" — 평균의 {rate(hi) / rate(allr):.0f}배" if rate(allr) else '', N)])
    set_paras(s[11], paras)
    t = s[12].text_frame.text
    set_text(s[12], re.sub(r'거래처 평균 [\d.]+%', f"거래처 평균 {f1(cur['revisitRate'])}%", t))
    fill_table(s[13].table, [[None, f"{r[1]}건", f"{r[2]}건", f"{f1(rate(r))}%", (f"{f1(prate(r))}%" if prate(r) is not None else '-'), (r[5].replace('-', '−') if r[5] != '-' else '-')] for r in ca['corrRows']])

    # ── 6. 트렌드 비교
    s = S[5]
    tc = d['trendCompare']; last7 = mr[-7:]
    replace_chart(s[1].chart, [f"{mnum(r[0])}월" for r in last7], [('PM건수', [r[2] for r in last7]), ('수리건수', [r[1] for r in last7])])
    def pct(a, b): return (a - b) / b * 100
    pmd, rpd = pct(tc['postPM'], tc['prePM']), pct(tc['postRP'], tc['preRP'])
    fill_table(s[3].table, [[None, comma(tc['prePM']), comma(tc['postPM']), sgn(pmd, 1, '%')],
                            [None, comma(tc['preRP']), comma(tc['postRP']), sgn(rpd, 1, '%')],
                            [None, f"{f1(tc['prePMRatio'])}%", f"{f1(tc['postPMRatio'])}%", sgn(tc['postPMRatio'] - tc['prePMRatio'], 1, '%p')]])
    npost = len(tc['postKeyLabels'])
    set_text(s[4], f"개편 전 = 2025.4~2026.3 · 개편 후 = 2026.4~{m}월 월평균")
    yoy = d.get('yoy')
    if yoy:
        a, c = yoy['prev'], yoy['cur']
        hdr = s[5].table
        set_cell(hdr.cell(0, 1), f"{yoy['prevLabel'][:4]}년 {mnum(yoy['prevLabel'])}월"); set_cell(hdr.cell(0, 2), f"{yoy['curLabel'][:4]}년 {m}월")
        fill_table(hdr, [[None, comma(a['repair']), comma(c['repair']), sgn(pct(c['repair'], a['repair']), 0, '%') if a['repair'] else '-'],
                         [None, comma(a['pm']), comma(c['pm']), sgn(pct(c['pm'], a['pm']), 0, '%') if a['pm'] else '-'],
                         [None, (f"{f1(a['ratio'])}%" if a['ratio'] is not None else '-'), f"{f1(c['ratio'])}%", sgn(c['ratio'] - a['ratio'], 1, '%p') if a['ratio'] is not None else '-']])
    else:
        SKIP.append('6장 「전년 대비」 표 — json 에 yoy 없음(엔진 최신판으로 다시 만들어야 함)')
    rank = sorted(((r[2], r[0]) for r in mr if r[0] in tc['postKeyLabels']), reverse=True)
    pos = [x[1] for x in rank].index(label) + 1 if label in [x[1] for x in rank] else None
    yoy_pm = f" · 전년 같은 달 대비 PM {sgn(pct(yoy['cur']['pm'], yoy['prev']['pm']), 0, '%')}" if yoy and yoy['prev']['pm'] else ''
    set_paras(s[7], [[('Insight', B)],
                     [(f"개편 후 {npost}개월 평균  ", B), ('PM ', N), (sgn(pmd, 1, '%'), B), (f" · 수리 {sgn(rpd, 1, '%')} · 예방정비 지수 ", N), (sgn(tc['postPMRatio'] - tc['prePMRatio'], 1, '%p'), B)],
                     [(f"{mo}  ", B), (f"PM {comma(pmv)}건(개편 후 {pos}위){yoy_pm}", N)]])

    # ── 7. 수도권
    s = S[6]
    metro = d['metro']; months6 = metro['months']
    by = {(r[0], r[1]): (r[2], r[3]) for r in metro['rows']}
    replace_chart(s[1].chart, [f"{mnum(l)}월" for l in months6],
                  [('PM건수', [by[(l, 'WEST')][0] + by[(l, 'EAST')][0] for l in months6]), ('수리건수', [by[(l, 'WEST')][1] + by[(l, 'EAST')][1] for l in months6])])
    tb = s[2].table
    for j, l in enumerate(months6): set_cell(tb.cell(0, 2 + j), f"{mnum(l)}월")
    fill_table(tb, [[None, None] + [str(by[(l, 'WEST')][1]) for l in months6], [None, None] + [str(by[(l, 'EAST')][1]) for l in months6],
                    [None, None] + [str(by[(l, 'WEST')][0]) for l in months6], [None, None] + [str(by[(l, 'EAST')][0]) for l in months6]])
    set_text(s[3], f"{mnum(months6[0])}~{m}월 = 개편 후 구간 · {m}월은 보고 대상 월 (개편 전 대비는 아래 표)")
    md = d.get('metroDelta')
    if md:
        def row(g): x = md[g]; return [sgn(x['visit'], 1, '%'), sgn(x['repair'], 1, '%'), sgn(x['pm'], 1, '%'), sgn(x['ratio'], 1, '%p')]
        fill_table(s[4].table, [[None] + row('FS WEST'), [None] + row('FS EAST'), [None] + row('WEST+EAST'), [None] + row('FSG')])
        set_text(s[5], re.sub(r'개편 후\(2026\.4~\d+\)', f'개편 후(2026.4~{m})', s[5].text_frame.text))
    else:
        SKIP.append('7장 「개편 전 대비」 팀별 표 — json 에 metroDelta 없음(엔진 최신판 필요)')
    REVIEW.append(f'7장 「수도권 포인트」 3줄 — 문장 새로 써야 함: "{s[7].text_frame.text[:40]}…"')

    # ── 8. 주요 고객·장애
    s = S[7]
    set_text(s[0], f'주요 고객 및 장애 분석 ({mo})'); set_text(s[1], f"{label[:4]}년 {m}월 단월 기준 — 당월 집중 발생 지점 확인")
    set_text(s[2], f'병원 TOP 5 ({mo} 활동 기준)'); set_text(s[3], f'반복 고장 장비 TOP 5 ({mo})'); set_text(s[5], f'주요 장애 요인 ({mo} Top 4) – 장비 구분 없음')
    set_text(s[11], s[11].text_frame.text.replace('서비스 케이스 건수', '활동(방문) 건수'))
    hm = d['hospMonth']
    if all('team' in h for h in hm):
        fill_table(s[12].table, [[str(i + 1), short_acct(h['account']), TEAM_SHORT.get(h['team'], h['team']), str(h['count'])] for i, h in enumerate(hm)])
    else:
        fill_table(s[12].table, [[str(i + 1), h['account'], None, str(h['count'])] for i, h in enumerate(hm)]); SKIP.append('8장 병원 TOP5 「담당」 열 — json 에 team 없음')
    rm = d['repeatMonth']
    fill_table(s[13].table, [[r['serial'], short_model(r.get('model') or r.get('modelAlt') or ''), short_acct(r['account'], True), (TEAM_SHORT.get(r.get('team'), r.get('team')) if r.get('team') else None), str(r['count'])] for r in rm])
    er = d['errMonth']['rows'][:4][::-1]
    short = lambda e: re.sub(r' is required| detected', '', e)[:24]
    replace_chart(s[6].chart, [short(r['error']) for r in er], [('건수', [r['count'] for r in er])])
    pk = d.get('peopleKpi')
    if pk and pk['top']:
        set_text(s[7], f'FSE KPI 종합점수 TOP 10 ({mo})')
        rows8 = []
        for x in pk['top']:
            dp = x['deltaPct']
            rows8.append([x['name'], x['unitLabel'], x['grade'], f1(x['score']), str(x['count']), (f"{arrow(dp)}{abs(dp):.1f}%" if dp is not None else '-')])
        fill_table(s[8].table, rows8)
        tb8 = s[8].table
        for i, x in enumerate(pk['top'], start=1):
            dp = x['deltaPct']
            if dp is not None and dp != 0: set_cell(tb8.cell(i, 5), f"{arrow(dp)}{abs(dp):.1f}%", color=(GREEN if dp > 0 else RED))
    else:
        SKIP.append('8장 「FSE KPI 종합점수 TOP 10」 표 — json 에 peopleKpi 없음(엔진 최신판 필요)')
    REVIEW.append('8장 「분석」 문장 — 병원·장애 요인 순위가 바뀌었으니 새로 써야 함')
    REVIEW.append('8장 표의 「담당」 — 확정 덱은 「FSR 담당자 리스트」의 FSR 이름까지 적었으나 엔진은 방문이 가장 많은 팀만 낸다. FSR 이름이 필요하면 담당자 리스트로 채울 것')

    # ── 10. 에러율 · 운영 건전성
    s = S[9]
    er12 = d['errRate']['months']
    replace_chart(s[2].chart, [yymm(r[0]) for r in er12], [(g, [r[1 + i] for r in er12]) for i, g in enumerate(d['errRate']['groups'])])
    h = d['health']; prevh = (d.get('prev') or {}).get('health')
    if h:
        set_text(s[1], re.sub(r'가동 [\d,]+대', f"가동 {comma(h['activeCount'])}대", s[1].text_frame.text))
        tb = s[5].table; set_cell(tb.cell(0, 1), mo); set_cell(tb.cell(0, 2), f'{pmo} 대비')
        def delta_cell(cur_v, prev_v, unit, up_good, nd=1):
            if prev_v is None or cur_v is None: return None, None
            dl = round(cur_v - prev_v, nd); good = (dl >= 0) if up_good else (dl <= 0)
            return f"{arrow(dl)}{abs(dl):.{nd}f}{unit}", (GREEN if good else RED) if dl != 0 else None
        cells = [('PM 커버리지율', h['pmCoverage'], '%', '%p', True, 1), ('점검 일정 고르기', h['pmCV'], '%', '%p', False, 0),
                 ('100대당 고장 발생', h['per100'], '건', '건', False, 1), ('반복 고장 장비 비율', h['chronicRatio'], '%', '%p', False, 1)]
        keyp = {'PM 커버리지율': 'pmCoverage', '점검 일정 고르기': 'pmCV', '100대당 고장 발생': 'per100', '반복 고장 장비 비율': 'chronicRatio'}
        for i, (nm, v, u, du, upg, nd) in enumerate(cells, start=1):
            set_cell(tb.cell(i, 1), (f"{v:.{nd}f}{u}" if nd else f"{v:.0f}{u}"))
            if prevh and prevh.get(keyp[nm]) is not None:
                txt, col = delta_cell(v, prevh[keyp[nm]], du, upg, nd); set_cell(tb.cell(i, 2), txt, color=col)
            else:
                SKIP.append(f'10장 건전성 「{nm}」 전월 대비 — json 에 prev.health 없음')
        cvs = [x for x in h['teamCV'].values() if x is not None]
        set_cell(tb.cell(5, 1), f"{min(cvs)}~{max(cvs)}%")
        tb2 = s[6].table
        for r in range(1, len(tb2.rows)):
            tk = next((k for k, v in TEAM_KO.items() if v == tb2.cell(r, 0).text.strip()), None)
            if tk and h['teamStat'].get(tk):
                st = h['teamStat'][tk]; set_cell(tb2.cell(r, 1), f"{comma(st['avg'])}건"); set_cell(tb2.cell(r, 2), f"{h['teamCV'][tk]}%"); set_cell(tb2.cell(r, 3), f"{st['min']} ~ {st['max']}")
    else:
        SKIP.append('10장 운영 건전성 — json 에 health 없음(Serial 필요)')
    REVIEW.append('10장 Insight·「업무량 편차 진단」 문장 — 새로 써야 함')
    SKIP.append('9장 FCA · 11장 만족도 · 12장 결과/Next Step — 별도 자료(FCA 발표 엑셀·VOC 월별보고)가 필요해 건드리지 않음')

    # 표지·12장 월 표기는 건드리지 않음 — 월이 바뀌면 사람이 확인
    prs.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('base'); ap.add_argument('json'); ap.add_argument('-o', '--out')
    a = ap.parse_args()
    data = json.load(open(a.json, encoding='utf8'))
    out = a.out or os.path.join(os.getcwd(), 'filled_' + os.path.basename(a.base))
    run(a.base, data, out)
    print('saved:', out)
    print('\n[건너뜀 — json 에 없거나 이 도구가 다루지 않음]'); [print(' -', x) for x in SKIP]
    print('\n[검토 필요 — 사람이 새로 써야 하는 문장]'); [print(' -', x) for x in REVIEW]


if __name__ == '__main__':
    main()
